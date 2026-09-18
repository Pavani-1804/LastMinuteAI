from docx import Document

def extract_docx_text(path):
    try:
        doc = Document(path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += "\n" + cell.text
        return text
    except Exception as e:
        print(f"Error parsing DOCX: {e}")
        return ""
